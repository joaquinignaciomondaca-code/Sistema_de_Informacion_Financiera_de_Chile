"""Reporte Financiero de Cooperativas de Ahorro y Crédito (CMF) → balance y resultados por cooperativa.

Fuente: https://www.cmfchile.cl/portal/estadisticas/626/w4-propertyvalue-28918.html
Formato cubierto: 2017-01 en adelante (hojas "Activos Cooperativas", "Pasivos Cooperativas",
"Estado Resultados Coop", "Margen Interes - Comisiones"; cifras en millones de pesos).
Antes de 2017 la CMF usaba otro plan de cuentas y otra planilla: no se mezcla.

Se usa un esquema fijo por hoja (orden de columnas numéricas) que se valida contra la cabecera
de la planilla y contra identidades contables. Cualquier desvío detiene el período (fail-closed):
  * la cabecera debe contener las frases clave del esquema;
  * cada fila de cooperativa debe tener exactamente N montos;
  * subtotales internos (p. ej. colocaciones = comerciales + personas) deben cuadrar;
  * Activos totales = Pasivos totales + Patrimonio (por cooperativa);
  * la suma de cooperativas debe igualar la fila "Total Cooperativas" (tolerancia de redondeo).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import unicodedata
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
INDEX = "https://www.cmfchile.cl/portal/estadisticas/626/w4-propertyvalue-28918.html"
FIRST_PERIOD = "2017-01"
OUT_DIR = ROOT / "docs" / "outputs" / "cooperativas" / "cmf_reporte_financiero"
MONTHS = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7,
          "agosto": 8, "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}

COOPERATIVAS = {
    # clave normalizada del nombre en la planilla → (RUT, nombre de fantasía)
    "coopeuch": ("82878900-7", "COOPEUCH"),
    "oriencoop": ("70010920-8", "ORIENCOOP"),
    "capual": ("84156800-1", "CAPUAL"),
    "ahorrocoop": ("81836800-3", "AHORROCOOP"),
    "detacoop": ("70017860-9", "DETACOOP"),
    "coonfia": ("70286300-7", "COONFIA"),
    "lautaro rosas": ("70286300-7", "COONFIA"),  # nombre anterior de la misma cooperativa
    "coocretal": ("70015260-K", "COOCRETAL"),
}


@dataclass(frozen=True)
class Sheet:
    key: str
    estado: str           # balance | resultados
    sheet_pattern: str    # regex sobre el nombre de hoja
    header_keys: tuple    # frases que deben aparecer en la cabecera
    concepts: tuple       # (codigo, glosa, nivel)
    identities: tuple     # (total_idx, (sumandos_idx...)) 0-based


ACTIVOS = Sheet(
    "activos", "balance", r"^activos\s+cooperativas", ("efectivo", "instrumentos", "colocaciones", "provisiones", "activos totales"),
    (
        ("efectivo_depositos_bancos", "Efectivo y depósitos en bancos", 1),
        ("instrumentos_no_derivados", "Instrumentos financieros no derivados — total", 1),
        ("instrumentos_negociacion", "Instrumentos para negociación", 2),
        ("instrumentos_disponibles_venta", "Instrumentos de inversión disponibles para la venta", 2),
        ("instrumentos_al_vencimiento", "Instrumentos de inversión hasta el vencimiento", 2),
        ("colocaciones", "Colocaciones — total", 1),
        ("colocaciones_comerciales", "Colocaciones comerciales (empresas)", 2),
        ("colocaciones_personas", "Colocaciones a personas — total", 2),
        ("colocaciones_consumo", "Colocaciones de consumo — total", 3),
        ("consumo_en_cuotas", "Consumo en cuotas", 4),
        ("consumo_descuento_planilla", "Consumo con descuento por planilla", 4),
        ("consumo_tarjetas_credito", "Consumo tarjetas de crédito", 4),
        ("consumo_otros", "Otros créditos de consumo", 4),
        ("colocaciones_vivienda", "Colocaciones para vivienda", 3),
        ("provisiones_constituidas", "Provisiones constituidas (colocaciones)", 1),
        ("activos_totales", "Activos totales", 0),
    ),
    ((1, (2, 3, 4)), (5, (6, 7)), (7, (8, 13)), (8, (9, 10, 11, 12))),
)
PASIVOS = Sheet(
    "pasivos", "balance", r"^pasivos\s+cooperativas", ("pasivos", "patrimonio", "dep", "capital pagado"),
    (
        ("pasivos_totales", "Pasivos totales", 0),
        ("depositos_captaciones", "Depósitos y captaciones — total", 1),
        ("depositos_vista", "Depósitos a la vista", 2),
        ("depositos_plazo", "Depósitos a plazo", 2),
        ("cuentas_ahorro", "Cuentas de ahorro", 2),
        ("prestamos_obtenidos", "Préstamos obtenidos — total", 1),
        ("prestamos_bancos_pais", "Préstamos de bancos del país", 2),
        ("prestamos_otras_instituciones", "Préstamos de otras instituciones", 2),
        ("instrumentos_deuda_emitidos", "Instrumentos de deuda emitidos (bonos)", 1),
        ("provisiones_adicionales", "Provisiones adicionales para colocaciones", 1),
        ("provisiones_contingentes", "Provisiones por riesgo de crédito contingente", 1),
        ("patrimonio_total", "Patrimonio — total", 0),
        ("capital_pagado", "Capital pagado", 1),
    ),
    ((1, (2, 3, 4)), (5, (6, 7))),
)
RESULTADOS = Sheet(
    "resultados", "resultados", r"^estado\s+resultados\s+coop", ("margen de intereses", "comisiones", "provisiones", "impuesto", "castigos"),
    (
        ("margen_intereses", "Margen de intereses", 1),
        ("comisiones_netas", "Comisiones netas", 1),
        ("resultado_operaciones_financieras", "Resultado neto de operaciones financieras", 1),
        ("recuperacion_castigados", "Recuperación de créditos castigados", 1),
        ("otros_ingresos_operacionales", "Otros ingresos operacionales netos", 1),
        ("resultado_operacional_bruto", "Resultado operacional bruto", 0),
        ("gasto_provisiones", "Gasto en provisiones — total", 1),
        ("provisiones_riesgo_colocaciones", "Provisiones por riesgo de crédito de colocaciones", 2),
        ("provisiones_adicionales_gasto", "Provisiones adicionales", 2),
        ("provisiones_riesgo_contingentes", "Provisiones por riesgo de créditos contingentes", 2),
        ("gastos_apoyo", "Gastos de apoyo", 1),
        ("resultado_operacional_neto", "Resultado operacional neto", 0),
        ("resultado_inversion_sociedades", "Resultado por inversión en sociedades", 1),
        ("resultado_antes_impuesto", "Resultado antes de impuesto", 0),
        ("impuesto", "Impuesto", 1),
        ("resultado_ejercicio", "Resultado del ejercicio", 0),
        ("resultado_propietarios", "Resultado atribuible a propietarios de la controladora", 1),
        ("resultado_no_controladoras", "Resultado atribuible a participaciones no controladoras", 1),
        ("castigos_ejercicio", "Castigos del ejercicio (informativo)", 9),
    ),
    ((5, (0, 1, 2, 3, 4)), (6, (7, 8, 9)), (11, (5, 6, 10)), (13, (11, 12)), (15, (13, 14)), (15, (16, 17))),
)
MARGEN = Sheet(
    "margen", "resultados", r"^margen\s+inter", ("margen de intereses", "ingresos por intereses", "gastos por intereses", "comisiones"),
    (
        ("margen_intereses_total", "Margen de intereses (desglose)", 1),
        ("ingresos_intereses_reajustes", "Ingresos por intereses y reajustes — total", 2),
        ("ingresos_colocaciones", "Ingresos de colocaciones — total", 3),
        ("ingresos_colocaciones_comerciales", "Ingresos de colocaciones comerciales", 4),
        ("ingresos_colocaciones_consumo", "Ingresos de colocaciones de consumo", 4),
        ("ingresos_colocaciones_vivienda", "Ingresos de colocaciones para vivienda", 4),
        ("ingresos_instrumentos_inversion", "Ingresos de instrumentos de inversión", 3),
        ("ingresos_otros", "Otros ingresos por intereses y reajustes", 3),
        ("gastos_intereses_reajustes", "Gastos por intereses y reajustes — total", 2),
        ("gastos_intereses", "Gastos por intereses", 3),
        ("gastos_reajustes", "Gastos por reajustes", 3),
        ("comisiones_netas_total", "Comisiones netas (desglose)", 1),
        ("ingresos_comisiones", "Ingresos por comisiones y servicios — total", 2),
        ("comisiones_seguros", "Comisiones por seguros", 3),
        ("comisiones_tarjetas", "Comisiones por tarjetas de crédito", 3),
        ("comisiones_otros_servicios", "Comisiones por otros servicios", 3),
        ("gastos_comisiones", "Gastos por comisiones", 2),
    ),
    ((0, (1, 8)), (1, (2, 6, 7)), (2, (3, 4, 5)), (8, (9, 10)), (11, (12, 16)), (12, (13, 14, 15))),
)
SHEETS = (ACTIVOS, PASIVOS, RESULTADOS, MARGEN)
IDENTITY_TOLERANCE = 2      # MM$: los subtotales publicados se redondean por separado
DECLARED_TOLERANCE = 5      # MM$: diferencias de la propia fuente CMF que se aceptan pero se declaran
TOTAL_TOLERANCE_PER_COOP = 1


def norm(text) -> str:
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", text).strip()


def coop_key(label) -> str | None:
    n = norm(re.sub(r"\(\d+\)", "", str(label)))
    for key in COOPERATIVAS:
        if n == key or n.startswith(key):
            return key
    return None


def is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def row_block(r: list) -> list:
    """Celdas tras la etiqueta hasta un hueco de ≥3 columnas vacías (área de trabajo a la derecha)."""
    start = next(i for i, v in enumerate(r) if v not in (None, "")) + 1
    out, gap = [], 0
    for v in r[start:]:
        if v in (None, ""):
            gap += 1
            if gap >= 3 and out:
                break
            continue
        gap = 0
        out.append(v)
    return out


def parse_sheet(spec: Sheet, rows: list[list]) -> tuple[dict, list[int]]:
    """Devuelve ({clave_coop: [montos]}, fila_total) o lanza ValueError con el motivo."""
    first = None
    for n, r in enumerate(rows):
        vals = [v for v in r if v not in (None, "")]
        if vals and isinstance(vals[0], str) and coop_key(vals[0]):
            first = n
            break
    if first is None:
        raise ValueError(f"{spec.key}: no se encontraron filas de cooperativas")
    header = norm(" ".join(str(v) for r in rows[:first] for v in r if isinstance(v, str)))
    missing = [k for k in spec.header_keys if k not in header]
    if missing:
        raise ValueError(f"{spec.key}: cabecera sin {missing}")
    width = len(spec.concepts)
    # Sólo columnas bajo la cabecera (algunas planillas traen columnas auxiliares ocultas a la derecha).
    last_col = max((i for r in rows[:first] for i, v in enumerate(r) if isinstance(v, str) and v.strip()), default=None)
    if last_col is not None:
        rows = rows[:first] + [list(r[:last_col + 1]) for r in rows[first:]]
    data: dict[str, list[int]] = {}
    total = None
    for r in rows[first:]:
        vals = [v for v in r if v not in (None, "")]
        if not vals or not isinstance(vals[0], str):
            continue
        label = vals[0]
        nums = [v for v in row_block(r) if is_number(v) or (isinstance(v, str) and v.strip() in ("---", "-"))]
        nums = [0 if isinstance(v, str) else v for v in nums]
        if norm(label).startswith("total cooperativas"):
            if len(nums) != width:
                raise ValueError(f"{spec.key}: fila total con {len(nums)} montos, se esperaban {width}")
            total = [round(v) for v in nums]
            break
        key = coop_key(label)
        if key is None:
            if re.search(r"coop", norm(label)):
                raise ValueError(f"{spec.key}: cooperativa no reconocida {label!r}")
            continue
        if len(nums) != width:
            raise ValueError(f"{spec.key}: {label} con {len(nums)} montos, se esperaban {width}")
        if key in data:
            raise ValueError(f"{spec.key}: {label} duplicada")
        data[key] = [round(v) for v in nums]
    if total is None:
        raise ValueError(f"{spec.key}: falta la fila 'Total Cooperativas'")
    return data, total


def check_identities(spec: Sheet, name: str, vals: list[int], declared: list) -> None:
    for total_idx, parts in spec.identities:
        s = sum(vals[i] for i in parts)
        diff = vals[total_idx] - s
        if IDENTITY_TOLERANCE < abs(diff) <= DECLARED_TOLERANCE:
            declared.append({"seccion": spec.key, "cooperativa": COOPERATIVAS[name][1], "concepto": spec.concepts[total_idx][0],
                             "publicado": vals[total_idx], "suma_componentes": s, "diferencia_mm_clp": diff})
        elif abs(diff) > DECLARED_TOLERANCE:
            raise ValueError(
                f"{spec.key}/{name}: {spec.concepts[total_idx][0]}={vals[total_idx]} ≠ suma {s} "
                f"de {[spec.concepts[i][0] for i in parts]}")


def validate_period(parsed: dict[str, tuple[dict, list[int]]]) -> dict:
    coops = set(parsed["activos"][0])
    declared: list[dict] = []
    for key, (data, _) in parsed.items():
        if set(data) != coops:
            raise ValueError(f"{key}: cooperativas {sorted(data)} ≠ activos {sorted(coops)}")
    for spec in SHEETS:
        data, total = parsed[spec.key]
        for name, vals in data.items():
            check_identities(spec, name, vals, declared)
        for i, (code, _, _) in enumerate(spec.concepts):
            s = sum(v[i] for v in data.values())
            if abs(s - total[i]) > TOTAL_TOLERANCE_PER_COOP * len(data):
                raise ValueError(f"{spec.key}: suma cooperativas {code}={s} ≠ Total Cooperativas {total[i]}")
    a, p = parsed["activos"][0], parsed["pasivos"][0]
    for name in coops:
        lhs, rhs = a[name][15], p[name][0] + p[name][11]
        if IDENTITY_TOLERANCE < abs(lhs - rhs) <= DECLARED_TOLERANCE:
            declared.append({"seccion": "balance", "cooperativa": COOPERATIVAS[name][1], "concepto": "activos = pasivos + patrimonio",
                             "publicado": lhs, "suma_componentes": rhs, "diferencia_mm_clp": lhs - rhs})
        elif abs(lhs - rhs) > DECLARED_TOLERANCE:
            raise ValueError(f"{name}: activos totales {lhs} ≠ pasivos + patrimonio {rhs}")
    r, m = parsed["resultados"][0], parsed["margen"][0]
    for name in coops:
        for ri, mi in ((0, 0), (1, 11)):
            if abs(r[name][ri] - m[name][mi]) > IDENTITY_TOLERANCE:
                raise ValueError(f"{name}: resultados/margen no calzan en {RESULTADOS.concepts[ri][0]}")
    return {"cooperativas": sorted(COOPERATIVAS[k][1] for k in coops), "n_cooperativas": len(coops),
            "diferencias_fuente_declaradas": declared}


def to_rows(period: str, parsed: dict, url: str, sha: str) -> list[dict]:
    y, mth = map(int, period.split("-"))
    import calendar
    fecha = f"{period}-{calendar.monthrange(y, mth)[1]:02d}"
    rows = []
    for spec in SHEETS:
        data, _ = parsed[spec.key]
        for key, vals in sorted(data.items()):
            rut, fantasia = COOPERATIVAS[key]
            for orden, ((code, glosa, nivel), v) in enumerate(zip(spec.concepts, vals), start=1):
                rows.append({
                    "id": f"{period}:{rut}:{spec.key}:{code}", "periodo": period, "fecha_corte": fecha,
                    "rut": rut, "cooperativa": fantasia, "estado": spec.estado, "seccion": spec.key,
                    "orden": orden, "codigo_concepto": code, "glosa": glosa, "nivel": nivel,
                    "monto_mm_clp": int(v),
                    "base_monto": "saldo al cierre" if spec.estado == "balance" else "acumulado del año a la fecha",
                    "fuente_url": url, "sha256_fuente": sha,
                })
    return rows


# ---------------------------------------------------------------- descarga (sólo en Actions)

def discover(page: str) -> dict[str, str]:
    """{periodo: url_planilla} desde el índice CMF (el recurso comparte id con el artículo)."""
    from bancos.scripts.inspect_cmf_bank_sample import AnchorParser
    p = AnchorParser(); p.feed(page); p.close()
    labels, resources = {}, {}
    for item in p.links:
        href, text = item["href"], re.sub(r"\s+", " ", item["text"]).strip()
        m = re.search(r"w4-article-(\d+)\.html", href)
        if m and "cooperativas" in norm(text):
            mm = re.search(r"(" + "|".join(MONTHS) + r")\s+(20\d\d)", norm(text))
            if mm:
                labels[m.group(1)] = f"{mm.group(2)}-{MONTHS[mm.group(1)]:02d}"
        r = re.search(r"articles-(\d+)_recurso_1\.xlsx?", href)
        if r:
            resources[r.group(1)] = href if href.startswith("http") else "https://www.cmfchile.cl/portal/estadisticas/626/" + href.lstrip("/")
    out: dict[str, str] = {}
    for aid, per in labels.items():
        if aid in resources:
            if per in out and out[per] != resources[aid]:
                raise ValueError(f"dos planillas para {per}")
            out[per] = resources[aid]
    return out


def read_workbook(blob: bytes, url: str) -> dict[str, list[list]]:
    if url.split("?")[0].lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(blob), read_only=True, data_only=True)
        return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    import xlrd
    bk = xlrd.open_workbook(file_contents=blob)
    return {sh.name: [sh.row_values(i) for i in range(sh.nrows)] for sh in bk.sheets()}


def parse_workbook(sheets: dict[str, list[list]]) -> dict:
    parsed = {}
    for spec in SHEETS:
        names = [n for n in sheets if re.search(spec.sheet_pattern, norm(n))]
        if len(names) != 1:
            raise ValueError(f"hoja {spec.key}: se esperaba 1, hay {names}")
        parsed[spec.key] = parse_sheet(spec, sheets[names[0]])
    return parsed


def run(publish: bool, desde: str, hasta: str | None) -> int:
    from bancos.scripts.inspect_cmf_bank_sample import MAX_FILE, MAX_PAGE, fetch
    import pyarrow as pa
    import pyarrow.parquet as pq
    sources = discover(fetch(INDEX, MAX_PAGE).decode("utf-8", "replace"))
    periods = sorted(p for p in sources if p >= desde and (hasta is None or p <= hasta))
    all_rows, report, failures = [], [], []
    for period in periods:
        url = sources[period]
        try:
            blob = fetch(url, MAX_FILE)
            parsed = parse_workbook(read_workbook(blob, url))
            info = validate_period(parsed)
            rows = to_rows(period, parsed, url.split("?")[0], hashlib.sha256(blob).hexdigest())
            all_rows.extend(rows)
            report.append({"periodo": period, "status": "passed", "registros": len(rows), **info})
        except Exception as exc:  # fail-closed por período, se reporta todo
            failures.append({"periodo": period, "error": f"{type(exc).__name__}: {exc}"[:400]})
            print(f"::warning title=Cooperativas {period}::{failures[-1]['error']}", flush=True)
    expected = [p for p in periods]
    print(f"::notice title=Cooperativas CMF::períodos={len(periods)} ok={len(report)} fallidos={len(failures)} registros={len(all_rows)}")
    if failures:
        print("::error title=Cooperativas CMF::Hay períodos que no pasan la validación; no se publica.")
        return 1
    if not publish:
        return 0
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    table = pa.Table.from_pylist(all_rows)
    pq.write_table(table, OUT_DIR / "estados.parquet", compression="zstd")
    (OUT_DIR / "validacion.json").write_text(json.dumps({
        "dataset": "cooperativas_cmf_reporte_financiero", "fuente": INDEX,
        "cobertura": {"desde": expected[0], "hasta": expected[-1], "periodos": len(expected)},
        "unidad": "millones de pesos (MM$)",
        "nota": "Balance: saldos al cierre. Resultados: acumulados del año a la fecha. Formato CMF 2017+; "
                "antes de 2017 la planilla y el plan de cuentas son distintos y no se incluyen.",
        "reglas": ["cabecera con frases clave", "N montos exactos por fila", "subtotales internos (±2 MM$; hasta ±5 MM$ se acepta y se declara)",
                   "activos = pasivos + patrimonio (misma tolerancia)", "suma cooperativas = Total Cooperativas",
                   "margen/comisiones iguales entre hojas"],
        "total_registros": len(all_rows), "periodos": report,
    }, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    return 0


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--publicar", action="store_true")
    ap.add_argument("--desde", default=FIRST_PERIOD)
    ap.add_argument("--hasta")
    a = ap.parse_args()
    sys.exit(run(a.publicar, max(a.desde, FIRST_PERIOD), a.hasta))


if __name__ == "__main__":
    main()
