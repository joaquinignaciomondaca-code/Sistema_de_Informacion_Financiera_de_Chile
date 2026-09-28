"""Sonda (sin publicación): estructura del Reporte Financiero de Cooperativas CMF.

Uso:  python -m cooperativas.scripts.probe_coop_layout index
      python -m cooperativas.scripts.probe_coop_layout xlsx 2026-07
"""
from __future__ import annotations

import re
import sys
from io import BytesIO

from bancos.scripts.inspect_cmf_bank_sample import MAX_FILE, MAX_PAGE, AnchorParser, fetch, find_source

INDEX = "https://www.cmfchile.cl/portal/estadisticas/626/w4-propertyvalue-28918.html"
CHUNK = 3900


def annotate(title: str, text: str, limit: int = 9) -> None:
    text = text.replace("%", "%25").replace("\r", "")
    parts = [text[i:i + CHUNK] for i in range(0, len(text), CHUNK)][:limit]
    for part in parts:
        print(f"::warning title={title}::" + part.replace("\n", "%0A"), flush=True)


def index() -> None:
    page = fetch(INDEX, MAX_PAGE).decode("utf-8", "replace")
    p = AnchorParser(); p.feed(page); p.close()
    rows = []
    for item in p.links:
        text = re.sub(r"\s+", " ", item["text"]).strip()
        href = item["href"]
        if "recurso" in href or "w4-article" in href:
            rows.append(f"{text[:80]} | {href[-60:]}")
    more = re.findall(r"[?&](?:page|pagina|p)=\d+", page)[:5]
    annotate("COOP index", f"links={len(rows)} paginacion={more}\n" + "\n".join(rows[:60]) + "\n...\n" + "\n".join(rows[-30:]))


def xlsx(period: str) -> None:
    from openpyxl import load_workbook
    url, label = find_source(INDEX, period, ".xlsx")
    wb = load_workbook(BytesIO(fetch(url, MAX_FILE)), read_only=True, data_only=True)
    out = [f"{period} {url} hojas={wb.sheetnames}"]
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        out.append(f"=== {ws.title} filas={len(rows)} cols={max((len(r) for r in rows), default=0)}")
        for n, r in enumerate(rows, 1):
            vals = [v for v in r if v is not None]
            if not vals:
                continue
            if n <= 12 or any(isinstance(v, str) and re.search(r"coopeuch|oriencoop|oriente|total|sistema", v, re.I) for v in vals[:3]):
                out.append(f"{n}: {[str(v)[:45] for v in vals[:24]]}")
    annotate(f"COOP {period}", "\n".join(out))


if __name__ == "__main__":
    {"index": lambda: index(), "xlsx": lambda: xlsx(sys.argv[2])}[sys.argv[1]]()
