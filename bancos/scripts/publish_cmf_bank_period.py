"""Valida y publica una sola partición mensual CMF B1/B2/R1.

El script es idempotente: un período presente en el manifiesto nunca se vuelve a
extraer ni se sobrescribe. Los importes se conservan al grano de cuenta, en sus
campos fuente originales; no se derivan métricas financieras.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import tempfile
from datetime import date
from pathlib import Path
from typing import Any

from bancos.scripts.extract_cmf_bank_lines import (
    extract_archive,
    reconcile_to_inspection,
)
from bancos.scripts.inspect_cmf_bank_sample import (
    MAX_FILE,
    XLSX_INDEX,
    ZIP_INDEX,
    fetch,
    find_source,
    period_label,
)

ROOT = Path(__file__).resolve().parents[2]
OUTPUT_ROOT = ROOT / "docs" / "outputs" / "bancos" / "cmf_b1_b2_r1"
PARTITION_MANIFEST = OUTPUT_ROOT / "manifest.json"
DATA_MANIFEST = ROOT / "data_manifest.json"
DATASET_ID = "bancos_cmf_lineas"
DATASET_NAME = "bancos.cmf_lineas_b1_b2_r1"
SEED_PERIOD = "2026-07"
BALANCE_SHEET = "Est. Situación Financ. Bancos"
RESULTS_SHEET = "Est. del Resultado Bancos"
SYSTEM_AGGREGATE_CODES = {"999"}
MINIMUM_INDIVIDUAL_INSTITUTIONS = 17


def previous_month(today: date) -> str:
    """Return the latest month that is fully closed."""
    year, month = today.year, today.month - 1
    if month == 0:
        year, month = year - 1, 12
    return f"{year:04d}-{month:02d}"


def next_month(period: str) -> str:
    year, month = map(int, period.split("-"))
    if month == 12:
        return f"{year + 1:04d}-01"
    return f"{year:04d}-{month + 1:02d}"


def next_unpublished_period(manifest: dict[str, Any], today: date) -> str | None:
    """Pick only the first missing month, so a late CMF release cannot create gaps."""
    published = {str(item["period"]) for item in manifest.get("periods", [])}
    candidate = SEED_PERIOD
    while candidate in published:
        candidate = next_month(candidate)
    return candidate if candidate <= previous_month(today) else None


def select_period(requested: str, manifest: dict[str, Any], today: date) -> str | None:
    """Reject manual backfills/skips; allow an already-published request as a no-op."""
    if not requested:
        return next_unpublished_period(manifest, today)
    period_label(requested)
    published = {str(item["period"]) for item in manifest.get("periods", [])}
    if requested in published:
        return requested
    next_period = next_unpublished_period(manifest, today)
    if requested != next_period:
        expected = next_period or "ninguno (no hay período cerrado pendiente)"
        raise ValueError(
            f"No se permite backfill ni saltar períodos: se solicitó {requested}; "
            f"el único período nuevo admisible es {expected}."
        )
    return requested


def is_period_published(period: str, manifest: dict[str, Any], output_root: Path = OUTPUT_ROOT) -> bool:
    """Treat either the ledger entry or an existing partition as a hard duplicate guard."""
    published = {str(item.get("period")): item for item in manifest.get("periods", [])}
    partition = output_root / period / "lineas.parquet"
    if period in published:
        recorded_path = str(published[period].get("file", ""))
        if not partition.is_file() or not recorded_path.endswith(f"/{period}/lineas.parquet"):
            raise RuntimeError(f"El período {period} figura publicado pero falta su Parquet")
        return True
    if partition.exists():
        raise RuntimeError(
            f"La partición {partition} ya existe, pero no está registrada en manifest.json; "
            "no se sobrescribe automáticamente."
        )
    return False


def read_workbook_rows(path: Path) -> list[dict[str, Any]]:
    """Read the two official CMF summary sheets once for all-bank tie-outs."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl es obligatorio para cotejar el Excel CMF") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    expected = {BALANCE_SHEET.casefold(), RESULTS_SHEET.casefold()}
    try:
        for sheet in workbook.worksheets:
            if sheet.title.strip().casefold() not in expected:
                continue
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = list(values)
                if any(value is not None for value in values):
                    rows.append({"sheet": sheet.title, "row_number": row_number, "values": values})
    finally:
        workbook.close()
    if not any(row["sheet"].strip().casefold() == BALANCE_SHEET.casefold() for row in rows):
        raise RuntimeError(f"Falta la hoja CMF requerida: {BALANCE_SHEET}")
    if not any(row["sheet"].strip().casefold() == RESULTS_SHEET.casefold() for row in rows):
        raise RuntimeError(f"Falta la hoja CMF requerida: {RESULTS_SHEET}")
    return rows


def validate_release(rows: list[dict], report: dict, workbook_rows: list[dict]) -> dict:
    """Fail closed unless coverage, row identity, B1 and R1 source tie-outs pass."""
    if report.get("file_count") != report.get("institution_count", 0) * 3:
        raise RuntimeError("Se esperaba exactamente un archivo B1, B2 y R1 por institución")
    counts = report.get("account_rows_by_family", {})
    if set(counts) != {"B1", "B2", "R1"} or any(int(counts[k]) <= 0 for k in ("B1", "B2", "R1")):
        raise RuntimeError(f"Cobertura de filas B1/B2/R1 vacía o incompleta: {counts}")
    expected_total = sum(int(value) for value in counts.values())
    if len(rows) != expected_total:
        raise RuntimeError(f"Conteo de filas no coincide: extracción={expected_total}, filas={len(rows)}")
    ids = [row.get("id") for row in rows]
    if any(not value for value in ids) or len(ids) != len(set(ids)):
        raise RuntimeError("IDs de registros vacíos o duplicados en la extracción")
    if any(row.get("periodo") != report.get("period") for row in rows):
        raise RuntimeError("Hay filas cuyo período difiere del período solicitado")
    if any(row.get("numero_fila_fuente") is None or row.get("codigo_cuenta") is None for row in rows):
        raise RuntimeError("Falta la cuenta o la posición física de alguna fila fuente")

    names_by_code: dict[str, str] = {}
    for file_report in report.get("file_reports", []):
        code = file_report["codigo_institucion"]
        name = file_report["nombre_institucion_fuente"]
        previous = names_by_code.setdefault(code, name)
        if previous != name:
            raise RuntimeError(f"El código {code} presenta nombres institucionales distintos")
    if set(names_by_code) != set(report.get("institution_codes", [])):
        raise RuntimeError("La cobertura institucional no coincide con los códigos de la extracción")

    inspection = {"workbook": {"bank_rows": workbook_rows}}
    institution_checks = []
    for code, name in sorted(names_by_code.items()):
        tieout = reconcile_to_inspection(rows, inspection, bank_code=code, bank_name=name)
        is_system_aggregate = code in SYSTEM_AGGREGATE_CODES
        b1_ok = tieout["b1_status"] == "passed"
        r1_matches = tieout["r1_exact_account_matches_in_xlsx"]
        core_r1 = {
            match["codigo_cuenta"] for match in r1_matches
            if match["codigo_cuenta"] in {"590000000", "594000000"}
        }
        r1_ok = bool(core_r1)
        # The CMF workbook may publish an aggregate under a different label.
        # Such a row is informational; independent named institutions are mandatory.
        if not is_system_aggregate and not b1_ok:
            raise RuntimeError(f"B1 TOTAL ACTIVOS no concilia con el XLSX para {code} {name}")
        if not is_system_aggregate and not r1_ok:
            raise RuntimeError(f"R1 no concilia una cuenta de resultado clave con el XLSX para {code} {name}")
        institution_checks.append({
            "codigo_institucion": code,
            "nombre_institucion_fuente": name,
            "es_agregado_sistema": is_system_aggregate,
            "b1_total_activos": "passed" if b1_ok else "not_comparable",
            "r1_cuentas_clave_coincidentes": sorted(core_r1),
            "r1_exact_account_match_count": len(r1_matches),
            "status": "passed" if (is_system_aggregate or (b1_ok and r1_ok)) else "failed",
        })

    named_checks = [item for item in institution_checks if not item["es_agregado_sistema"]]
    if len(named_checks) < MINIMUM_INDIVIDUAL_INSTITUTIONS:
        raise RuntimeError(
            f"Sólo se pudieron cotejar {len(named_checks)} instituciones individuales; "
            f"mínimo requerido {MINIMUM_INDIVIDUAL_INSTITUTIONS}"
        )
    if any(item["status"] != "passed" for item in named_checks):
        raise RuntimeError("Falló al menos un cotejo de institución individual")

    return {
        "status": "passed",
        "period": report["period"],
        "institution_count": report["institution_count"],
        "individual_institutions_tied_out": len(named_checks),
        "file_count": report["file_count"],
        "account_rows_by_family": counts,
        "total_account_rows": len(rows),
        "b2_validation": "estructura, forma, códigos y cobertura CMF; el XLSX de resumen no publica B2 individual",
        "amounts_policy": "se conserva el dato de origen sin renombrar los cuatro campos B1/B2 ni calcular métricas",
        "institution_checks": institution_checks,
    }


def atomic_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, delete=False) as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
        temp = Path(handle.name)
    os.replace(temp, path)


def publish_period(period: str, dry_run: bool = False) -> dict:
    period_label(period)
    if period > previous_month(date.today()):
        raise ValueError(f"No se publica un período futuro o aún abierto: {period}")
    manifest = json.loads(PARTITION_MANIFEST.read_text(encoding="utf-8")) if PARTITION_MANIFEST.exists() else {
        "dataset": DATASET_ID, "periods": [], "files": [], "total_records": 0,
    }
    if is_period_published(period, manifest):
        print(f"Período {period} ya publicado; no se descarga ni se repite.")
        return {"status": "already_published", "period": period, "published_changed": False}

    zip_url, _ = find_source(ZIP_INDEX, period, ".zip")
    xlsx_url, _ = find_source(XLSX_INDEX, period, ".xlsx")
    zip_data = fetch(zip_url, MAX_FILE)
    xlsx_data = fetch(xlsx_url, MAX_FILE)
    zip_hash = hashlib.sha256(zip_data).hexdigest()
    xlsx_hash = hashlib.sha256(xlsx_data).hexdigest()
    rows, extraction = extract_archive(zip_data, period, zip_url)
    workbook_rows = read_workbook_rows_from_bytes(xlsx_data)
    validation = validate_release(rows, extraction, workbook_rows)

    for row in rows:
        row["fecha_corte"] = date.fromisoformat(row["fecha_corte"])
        row["fuente_url_zip"] = zip_url
        row["sha256_zip"] = zip_hash
        row["fuente_url_xlsx_cotejo"] = xlsx_url
        row["sha256_xlsx_cotejo"] = xlsx_hash
    period_dir = OUTPUT_ROOT / period
    parquet_path = period_dir / "lineas.parquet"
    report_path = period_dir / "validacion.json"
    if period_dir.exists():
        raise RuntimeError(f"La carpeta {period_dir} ya existe sin registro de período; no se sobrescribe")

    if dry_run:
        return {
            "status": "validated_not_published", "period": period,
            "validation": validation, "published_changed": False,
        }

    extraction["validation"]["published"] = True
    extraction["status"] = "validated_and_published"
    try:
        import pyarrow as pa
        import pyarrow.parquet as pq
    except ImportError as exc:
        raise RuntimeError("pyarrow es obligatorio para crear las particiones Parquet") from exc
    table = pa.Table.from_pylist(rows)
    period_dir.mkdir(parents=True, exist_ok=False)
    temp_parquet = period_dir / "lineas.parquet.tmp"
    old_partition_manifest = PARTITION_MANIFEST.read_bytes() if PARTITION_MANIFEST.exists() else None
    old_data_manifest = DATA_MANIFEST.read_bytes()
    try:
        pq.write_table(table, temp_parquet, compression="zstd", version="2.6")
        # Verify the generated Parquet round-trips and preserves the expected row count.
        checked = pq.read_table(temp_parquet)
        if checked.num_rows != len(rows) or checked.column_names != table.column_names:
            raise RuntimeError("Parquet temporal no conserva filas/columnas de la extracción")
        os.replace(temp_parquet, parquet_path)
        source_record = {
            "period": period,
            "file": parquet_path.relative_to(ROOT / "docs").as_posix(),
            "validation_file": report_path.relative_to(ROOT / "docs").as_posix(),
            "records": len(rows),
            "sha256_zip": zip_hash,
            "sha256_xlsx": xlsx_hash,
            "zip_url": zip_url,
            "xlsx_url": xlsx_url,
            "validated_at": date.today().isoformat(),
        }
        new_periods = sorted([*manifest.get("periods", []), source_record], key=lambda item: item["period"])
        if len({item["period"] for item in new_periods}) != len(new_periods):
            raise RuntimeError(f"El manifiesto ya contiene el período {period}")
        new_manifest = {
            "dataset": DATASET_ID,
            "description": "Líneas contables B1/B2/R1 CMF al grano fuente; importes sin reinterpretar.",
            "updated_at": date.today().isoformat(),
            "periods": new_periods,
            "files": [item["file"] for item in new_periods],
            "total_records": sum(int(item["records"]) for item in new_periods),
        }
        atomic_json(report_path, {
            "period": period,
            "status": "passed",
            "sources": {"zip_url": zip_url, "zip_sha256": zip_hash, "xlsx_url": xlsx_url, "xlsx_sha256": xlsx_hash},
            "extraction": extraction,
            "validation": validation,
            "published_at": date.today().isoformat(),
        })
        atomic_json(PARTITION_MANIFEST, new_manifest)
        update_data_manifest(new_manifest, source_record)
    except Exception:
        # Roll back both ledgers if any final write fails; never leave a period
        # marked as published without its validated Parquet/report.
        if old_partition_manifest is None:
            PARTITION_MANIFEST.unlink(missing_ok=True)
        else:
            PARTITION_MANIFEST.write_bytes(old_partition_manifest)
        DATA_MANIFEST.write_bytes(old_data_manifest)
        for path in (temp_parquet, parquet_path, report_path):
            path.unlink(missing_ok=True)
        try:
            period_dir.rmdir()
        except OSError:
            pass
        raise

    result = {
        "status": "published", "period": period, "published_changed": True,
        "rows": len(rows), "partition": source_record["file"], "validation": validation,
    }
    write_github_output(result)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return result


def read_workbook_rows_from_bytes(blob: bytes) -> list[dict[str, Any]]:
    with tempfile.NamedTemporaryFile(suffix=".xlsx") as handle:
        handle.write(blob)
        handle.flush()
        return read_workbook_rows(Path(handle.name))


def update_data_manifest(period_manifest: dict, latest: dict) -> None:
    catalog = json.loads(DATA_MANIFEST.read_text(encoding="utf-8"))
    tables = catalog.setdefault("tables", [])
    entry = next((item for item in tables if item.get("id") == DATASET_ID), None)
    total = int(period_manifest["total_records"])
    start, end = period_manifest["periods"][0]["period"], period_manifest["periods"][-1]["period"]
    if entry is None:
        entry = {
            "id": DATASET_ID,
            "name": DATASET_NAME,
            "view_name": DATASET_ID,
            "sector": "bancos",
            "sector_label": "Banca Comercial",
            "norma": "CMF — Manual del Sistema de Información para bancos (B1/MB1, B2/MB2, R1/MR1)",
            "corte": f"{start} a {end}",
            "frescura": "Actualización mensual validada contra reporte CMF",
            "modo": "Automático con gate de publicación",
            "ultima_actualizacion": date.today().isoformat(),
            "file_parquet": latest["file"],
            "files_manifest": latest["file"].rsplit("/", 2)[0] + "/manifest.json",
            "registros_reales": total,
            "descripcion": "Filas fuente de B1 consolidado, B2 individual y R1 consolidado. Campos monetarios originales sin reinterpretación; revisar tipo_linea antes de agregar cuentas jerárquicas.",
            "origen": latest["zip_url"],
        }
        tables.append(entry)
        catalog["total_tables"] = int(catalog.get("total_tables", len(tables) - 1)) + 1
    else:
        entry.update({
            "corte": f"{start} a {end}",
            "ultima_actualizacion": date.today().isoformat(),
            "file_parquet": latest["file"],
            "files_manifest": latest["file"].rsplit("/", 2)[0] + "/manifest.json",
            "registros_reales": total,
            "origen": latest["zip_url"],
        })
    catalog["total_records"] = int(catalog.get("total_records", 0)) + int(latest["records"])
    catalog["updated_at"] = date.today().isoformat()
    atomic_json(DATA_MANIFEST, catalog)


def write_github_output(result: dict) -> None:
    output = os.environ.get("GITHUB_OUTPUT")
    if output:
        with open(output, "a", encoding="utf-8") as handle:
            handle.write(f"published_changed={str(result.get('published_changed', False)).lower()}\n")
            handle.write(f"period={result.get('period', '')}\n")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--period", default="", help="Período YYYY-MM; por defecto siguiente mes pendiente desde 2026-07")
    parser.add_argument("--dry-run", action="store_true", help="Validar sin escribir archivos de publicación")
    args = parser.parse_args()
    try:
        manifest = json.loads(PARTITION_MANIFEST.read_text(encoding="utf-8")) if PARTITION_MANIFEST.exists() else {"periods": []}
        period = select_period(args.period, manifest, date.today())
        if period is None:
            result = {"status": "up_to_date", "period": "", "published_changed": False}
            print("No hay un período CMF cerrado pendiente; no se modifica la web.")
        else:
            result = publish_period(period, dry_run=args.dry_run)
        write_github_output(result)
    except Exception as exc:
        message = f"Publicación CMF bancaria detenida de forma segura: {type(exc).__name__}: {exc}"
        print(message, file=sys.stderr, flush=True)
        print("::error title=Gate CMF B1/B2/R1::" + message.replace("%", "%25").replace("\r", "%0D"), flush=True)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
