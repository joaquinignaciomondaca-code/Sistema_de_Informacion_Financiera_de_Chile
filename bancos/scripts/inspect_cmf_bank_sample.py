"""Inspect an official monthly CMF bank ZIP and workbook without publishing data.

The script downloads only the requested CMF period, reads the ZIP in memory (it
never extracts untrusted member paths), and writes a bounded structural report
plus the public source files for manual review. It does not touch docs/outputs.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from collections import Counter
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path

HOST = "www.cmfchile.cl"
ZIP_INDEX = "https://www.cmfchile.cl/portal/estadisticas/626/w4-propertyvalue-30250.html"
XLSX_INDEX = "https://www.cmfchile.cl/portal/estadisticas/626/w4-propertyvalue-28911.html"
# The CMF July-2026 resource anchors have no accessible text in raw HTML; these
# exact URLs were resolved from the official index page and are a bounded
# fallback for the default one-off sample.
KNOWN_SOURCES = {
    ("2026-07", ".zip"): "https://www.cmfchile.cl/portal/estadisticas/626/articles-113067_recurso_1.zip?ts=1787944720",
    ("2026-07", ".xlsx"): "https://www.cmfchile.cl/portal/estadisticas/626/articles-113057_recurso_1.xlsx?ts=1787943075",
}
MAX_PAGE = 3_000_000
MAX_FILE = 15_000_000
MAX_ZIP_TOTAL = 80_000_000
MEMBER_RE = re.compile(r"(?:^|/)b([12])(20\d{2})(\d{2})(\d{3})\.txt$", re.I)
MONTHS = {
    1: "enero", 2: "febrero", 3: "marzo", 4: "abril", 5: "mayo", 6: "junio",
    7: "julio", 8: "agosto", 9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre",
}


class AnchorParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.links: list[dict[str, str]] = []
        self.current: dict[str, str] | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag == "a":
            if self.current is not None:
                self.links.append(self.current)
            attrs_dict = dict(attrs)
            self.current = {"href": attrs_dict.get("href", ""), "text": ""}

    def handle_data(self, data: str) -> None:
        if self.current is not None:
            self.current["text"] += data

    def handle_endtag(self, tag: str) -> None:
        if tag == "a" and self.current is not None:
            self.links.append(self.current)
            self.current = None

    def close(self) -> None:
        super().close()
        if self.current is not None:
            self.links.append(self.current)
            self.current = None


def fetch(url: str, limit: int) -> bytes:
    parsed = urllib.parse.urlsplit(url)
    if parsed.scheme != "https" or parsed.hostname != HOST:
        raise ValueError(f"Refusing non-CMF source: {parsed.hostname}")
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "MonitorFinancieroChile-CMF-sample-inspection/1.0",
            "Accept": "text/html,application/zip,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
        },
    )
    for attempt in range(3):
        try:
            with urllib.request.urlopen(request, timeout=40) as response:
                final = urllib.parse.urlsplit(response.geturl())
                if response.status != 200 or final.hostname != HOST or final.scheme != "https":
                    raise ValueError("Unexpected CMF response or redirect")
                data = response.read(limit + 1)
            if len(data) > limit:
                raise ValueError(f"CMF response exceeds {limit} bytes")
            return data
        except urllib.error.HTTPError as exc:
            if exc.code in (429, 500, 502, 503, 504) and attempt < 2:
                time.sleep(2 ** attempt)
                continue
            raise RuntimeError(f"CMF returned HTTP {exc.code} for {parsed.path}") from None
        except (urllib.error.URLError, TimeoutError, ConnectionError) as exc:
            if attempt < 2:
                time.sleep(2 ** attempt)
                continue
            raise RuntimeError(f"CMF connection failed ({type(exc).__name__}) for {parsed.path}") from None
    raise RuntimeError("CMF download retries exhausted")


def period_label(period: str) -> str:
    if not re.fullmatch(r"20\d{2}-(?:0[1-9]|1[0-2])", period):
        raise ValueError("Period must be YYYY-MM")
    year, month = map(int, period.split("-"))
    return f"{MONTHS[month]} {year}"


def find_source(index_url: str, period: str, suffix: str) -> tuple[str, str]:
    label = period_label(period).casefold()
    page = fetch(index_url, MAX_PAGE).decode("utf-8", errors="replace")
    parser = AnchorParser()
    parser.feed(page)
    parser.close()
    matches: dict[str, str] = {}
    for item in parser.links:
        href = urllib.parse.urljoin(index_url, item["href"].strip())
        parsed = urllib.parse.urlsplit(href)
        text = re.sub(r"\s+", " ", item["text"]).strip().casefold()
        if (
            parsed.scheme == "https"
            and parsed.hostname == HOST
            and parsed.path.startswith("/portal/estadisticas/626/")
            and parsed.path.lower().endswith(suffix)
            and label in text
        ):
            matches[href] = text
    if len(matches) != 1 and (period, suffix) in KNOWN_SOURCES:
        url = KNOWN_SOURCES[(period, suffix)]
        parsed = urllib.parse.urlsplit(url)
        if parsed.scheme == "https" and parsed.hostname == HOST and parsed.path.startswith("/portal/estadisticas/626/") and parsed.path.lower().endswith(suffix):
            print("::notice title=CMF sample source fallback::" + json.dumps({"period": period, "suffix": suffix, "url": url}, ensure_ascii=False))
            return url, f"official-index-url-fallback:{period}{suffix}"
    if len(matches) != 1:
        candidates = [
            {"url": urllib.parse.urljoin(index_url, item["href"].strip()), "text": re.sub(r"\s+", " ", item["text"]).strip()[:120]}
            for item in parser.links
            if urllib.parse.urlsplit(urllib.parse.urljoin(index_url, item["href"].strip())).path.lower().endswith(suffix)
        ]
        raise RuntimeError(
            f"Expected one CMF {suffix} link for {period}; found {len(matches)}; "
            f"candidate_links={json.dumps(candidates, ensure_ascii=False)}"
        )
    url, text = next(iter(matches.items()))
    print("::notice title=CMF sample source::" + json.dumps({"period": period, "suffix": suffix, "url": url}, ensure_ascii=False))
    return url, text


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def safe_txt_summary(zf: zipfile.ZipFile, info: zipfile.ZipInfo) -> dict:
    data = zf.read(info)
    if all(byte < 0x80 for byte in data):
        decoded = data.decode("ascii")
        encoding = "ascii-compatible (UTF-8/Latin-1 indistinguishable)"
    else:
        try:
            decoded = data.decode("utf-8-sig")
            encoding = "utf-8-sig"
        except UnicodeDecodeError:
            decoded = data.decode("latin-1")
            encoding = "latin-1"
    lines = [line for line in decoded.splitlines() if line.strip()]
    widths = Counter(len(line.split("\t")) for line in lines)
    return {
        "name": info.filename,
        "compressed_bytes": info.compress_size,
        "uncompressed_bytes": info.file_size,
        "encoding": encoding,
        "nonempty_lines": len(lines),
        "tab_field_counts": dict(sorted(widths.items())),
        "preview_lines": lines[:5],
    }


def inspect_zip(blob: bytes, period: str, bank_code: str) -> dict:
    try:
        zf = zipfile.ZipFile(__import__("io").BytesIO(blob))
    except zipfile.BadZipFile:
        raise ValueError("Downloaded CMF bank archive is not a valid ZIP") from None
    with zf:
        infos = [item for item in zf.infolist() if not item.is_dir()]
        if sum(item.file_size for item in infos) > MAX_ZIP_TOTAL:
            raise ValueError("ZIP expands beyond the safe inspection limit")
        if any(item.file_size > MAX_FILE for item in infos):
            raise ValueError("ZIP contains an oversized member")
        members = []
        bank_files = {"1": [], "2": []}
        for info in infos:
            match = MEMBER_RE.search(info.filename)
            if match and f"{match.group(2)}-{match.group(3)}" == period:
                members.append(info.filename)
                if match.group(4) == bank_code:
                    bank_files[match.group(1)].append(info)
        if not members:
            raise RuntimeError(f"No B1/B2 TXT members found for {period}")
        missing = [kind for kind, files in bank_files.items() if len(files) != 1]
        if missing:
            raise RuntimeError(
                f"Expected one B{','.join(missing)} file for bank {bank_code}/{period}; "
                f"found {[len(bank_files[k]) for k in missing]}"
            )
        return {
            "archive_member_count": len(infos),
            "matching_period_b1_b2_count": len(members),
            "matching_period_b1_b2_names": sorted(members),
            "bank_code": bank_code,
            "bank_files": {
                f"B{kind}": safe_txt_summary(zf, files[0])
                for kind, files in sorted(bank_files.items())
            },
        }


def norm(value: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def inspect_workbook(path: Path, bank_name: str) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to inspect the CMF workbook") from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    matches = []
    needle = norm(bank_name)
    try:
        for ws in wb.worksheets:
            dimensions = ws.calculate_dimension()
            count_nonempty = 0
            found = []
            for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
                values = list(row)
                if any(value is not None for value in values):
                    count_nonempty += 1
                if any(needle in norm(value) for value in values if value is not None):
                    found.append({"row_number": row_number, "values": values})
                    if len(found) >= 5:
                        break
            sheets.append({
                "name": ws.title,
                "reported_dimension": dimensions,
                "nonempty_rows_scanned_until_match_limit": count_nonempty,
                "contains_bank_name": bool(found),
            })
            for found_row in found:
                matches.append({"sheet": ws.title, **found_row})
        if not matches:
            raise RuntimeError(f"Bank name {bank_name!r} not found in the CMF workbook")
        return {"sheet_count": len(sheets), "sheets": sheets, "bank_name": bank_name, "bank_rows": matches}
    finally:
        wb.close()


def run(period: str, bank_code: str, bank_name: str, output: Path) -> dict:
    zip_url, zip_label = find_source(ZIP_INDEX, period, ".zip")
    xlsx_url, xlsx_label = find_source(XLSX_INDEX, period, ".xlsx")
    zip_data = fetch(zip_url, MAX_FILE)
    xlsx_data = fetch(xlsx_url, MAX_FILE)
    output.mkdir(parents=True, exist_ok=True)
    zip_path = output / f"cmf_bancos_{period}.zip"
    xlsx_path = output / f"cmf_reporte_bancario_{period}.xlsx"
    zip_path.write_bytes(zip_data)
    xlsx_path.write_bytes(xlsx_data)
    zip_inspection = inspect_zip(zip_data, period, bank_code)
    workbook_inspection = inspect_workbook(xlsx_path, bank_name)
    report = {
        "status": "inspected_not_published",
        "period": period,
        "bank_code": bank_code,
        "bank_name": bank_name,
        "checked_at_utc": datetime.now(timezone.utc).isoformat(),
        "sources": {
            "zip": {"url": zip_url, "index_label": zip_label, "bytes": len(zip_data), "sha256": sha256(zip_data)},
            "xlsx": {"url": xlsx_url, "index_label": xlsx_label, "bytes": len(xlsx_data), "sha256": sha256(xlsx_data)},
        },
        "zip": zip_inspection,
        "workbook": workbook_inspection,
        "limitations": [
            "The workbook is the CMF monthly system report with institution-level rows; it is not a separate bank-issued workbook.",
            "No normalization, schema mapping, financial reconciliation, database write, or web publication was performed.",
        ],
    }
    (output / "inspection.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    compact = {
        "period": period,
        "bank": f"{bank_code} {bank_name}",
        "zip_members": zip_inspection["archive_member_count"],
        "txt_names": zip_inspection["matching_period_b1_b2_names"],
        "b1_rows": zip_inspection["bank_files"]["B1"]["nonempty_lines"],
        "b1_fields": zip_inspection["bank_files"]["B1"]["tab_field_counts"],
        "b2_rows": zip_inspection["bank_files"]["B2"]["nonempty_lines"],
        "b2_fields": zip_inspection["bank_files"]["B2"]["tab_field_counts"],
        "b1_preview": zip_inspection["bank_files"]["B1"]["preview_lines"],
        "b2_preview": zip_inspection["bank_files"]["B2"]["preview_lines"],
        "xlsx_sheet_names": [sheet["name"] for sheet in workbook_inspection["sheets"]],
        "xlsx_bank_rows": len(workbook_inspection["bank_rows"]),
        "xlsx_bank_rows_preview": workbook_inspection["bank_rows"][:6],
        "published": False,
    }
    print("::notice title=CMF bank sample inspection::" + json.dumps(compact, ensure_ascii=False, separators=(",", ":")))
    print(json.dumps(compact, ensure_ascii=False))
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--period", default="2026-07", help="CMF period YYYY-MM (default: 2026-07)")
    parser.add_argument("--bank-code", default="001", help="Three-digit CMF bank code (default: 001)")
    parser.add_argument("--bank-name", default="Banco de Chile", help="Name to locate in the CMF workbook")
    parser.add_argument("--output", type=Path, default=Path(".local-data/review/bancos/sample_inspection"))
    args = parser.parse_args()
    try:
        run(args.period, args.bank_code, args.bank_name, args.output)
    except (OSError, ValueError, RuntimeError, zipfile.BadZipFile) as exc:
        message = f"CMF sample inspection failed safely: {exc}"
        print(message, file=sys.stderr)
        print("::error title=CMF sample inspection failed::" + message.replace("\\n", " ").replace("%", "%25").replace("\r", "%0D"))
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
