"""Publica el balance de patrimonios separados entregado por NotebookLM.

Lee el xlsx subido en DATA_NUEVA y escribe tres tablas. No recalcula montos
ni descarta filas. El glosario dentro del libro dice 11.086 filas y 497
documentos; las hojas traen otra cantidad. Se publica lo que está en las hojas.
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "DATA_NUEVA" / "FSB_Patrimonio_Separado (1).xlsx"
OUT = ROOT / "docs" / "outputs" / "securitizadoras"


def periodo(value):
    if value is None or value == "":
        return None
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text.zfill(6) if text.isdigit() else text


def texto(value):
    if value is None:
        return None
    return str(value)


def numero(value):
    if value is None or value == "":
        return None
    return value


def filas(ws):
    rows = ws.iter_rows(values_only=True)
    header = [str(c) if c is not None else "" for c in next(rows)]
    for row in rows:
        if any(c is not None and c != "" for c in row):
            yield header, row


def publicar_cuentas(ws):
    out = []
    for i, (header, row) in enumerate(filas(ws), start=1):
        raw = dict(zip(header, row))
        out.append(
            {
                "id_linea": f"{raw.get('documento_id')}_{i}",
                "documento_id": raw.get("documento_id"),
                "archivo": texto(raw.get("archivo")),
                "entidad_rut": texto(raw.get("entidad_rut")),
                "entidad_nombre": texto(raw.get("entidad_nombre")),
                "patrimonio_codigo": texto(raw.get("patrimonio_codigo")),
                "periodo": periodo(raw.get("periodo")),
                "anio": raw.get("anio"),
                "mes": raw.get("mes"),
                "categoria": texto(raw.get("categoria")),
                "asiento": texto(raw.get("asiento")),
                "monto": numero(raw.get("monto")),
                "categoria_fsb": texto(raw.get("categoria_fsb")),
                "monto_normalizado": numero(raw.get("monto_normalizado")),
            }
        )
    return out


def publicar_vehiculos(ws):
    out = []
    for header, row in filas(ws):
        raw = dict(zip(header, row))
        out.append(
            {
                "documento_id": raw.get("documento_id"),
                "archivo": texto(raw.get("archivo")),
                "entidad_rut": texto(raw.get("entidad_rut")),
                "entidad_nombre": texto(raw.get("entidad_nombre")),
                "patrimonio_codigo": texto(raw.get("patrimonio_codigo")),
                "periodo": periodo(raw.get("periodo")),
                "anio": raw.get("anio"),
                "mes": raw.get("mes"),
                "total_financial_assets": numero(raw.get("total_financial_assets")),
                "loans": numero(raw.get("loans")),
                "short_term_assets": numero(raw.get("short_term_assets")),
                "short_term_liabilities": numero(raw.get("short_term_liabilities")),
                "long_term_liabilities": numero(raw.get("long_term_liabilities")),
                "equity": numero(raw.get("equity")),
                "ci2": numero(raw.get("CI2_credit_intermediation")),
                "mt2": numero(raw.get("MT2_maturity_transformation")),
                "l5": numero(raw.get("L5_leverage")),
            }
        )
    return out


def publicar_periodos(ws):
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c) for c in rows[1]]
    out = []
    for row in rows[2:]:
        raw = dict(zip(header, row))
        item = {"periodo": periodo(raw.get("periodo"))}
        for key, value in raw.items():
            if key == "periodo":
                continue
            item[key] = numero(value)
        out.append(item)
    return out


def escribir(nombre, filas_out):
    json_path = OUT / f"{nombre}.json"
    parquet_path = OUT / f"{nombre}.parquet"
    json_path.write_text(
        json.dumps(filas_out, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    table = pa.Table.from_pylist(filas_out)
    pq.write_table(table, parquet_path)
    print(nombre, len(filas_out), json_path.stat().st_size)


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    escribir("patrimonios_separados_notebooklm_cuentas", publicar_cuentas(wb["Detalle_de_Cuentas"]))
    escribir("patrimonios_separados_notebooklm_vehiculos", publicar_vehiculos(wb["Detalle_por_patrimonio"]))
    escribir("patrimonios_separados_notebooklm_periodos", publicar_periodos(wb["Agregado_Risk_Metrics_EF5"]))
    wb.close()


if __name__ == "__main__":
    main()
